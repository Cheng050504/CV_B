"""Unit tests for the run-preserving DOCX renderer.

The trickiest part of this project is ``replace_in_runs_preserve``:
it has to correctly handle placeholders that Word has silently split
across multiple ``<w:r>`` runs while preserving the formatting on
every surviving run. These tests pin that behavior.
"""

from __future__ import annotations

import io
from docx import Document

from resume_builder.services.docx_renderer import (
    compose_state_country,
    pack_repeating,
    replace_in_runs_preserve,
    sanitize_filename,
)


def _new_doc_with_paragraph(runs):
    """Build a fresh doc and return the paragraph that holds `runs`.

    `runs` is a list of (text, bold) tuples so we can later assert that
    formatting was preserved on runs that weren't rewritten.
    """
    doc = Document()
    p = doc.add_paragraph()
    for text, bold in runs:
        r = p.add_run(text)
        r.bold = bold
    return doc, p


# ------------------- tiny helper coverage -------------------

def test_sanitize_filename_strips_unsafe_chars():
    assert sanitize_filename("Jane Doe / CV?") == "Jane_Doe_CV"


def test_sanitize_filename_fallback():
    assert sanitize_filename("") == "Document"
    assert sanitize_filename("///") == "Document"


def test_compose_state_country_us_strips_country():
    assert compose_state_country("NY", "USA") == "NY"
    assert compose_state_country("CA", "United States") == "CA"


def test_compose_state_country_non_us_joins():
    assert compose_state_country("Ontario", "Canada") == "Ontario, Canada"
    assert compose_state_country("", "UK") == "UK"
    assert compose_state_country("", "") == ""


def test_pack_repeating_dict_interface():
    form = {
        "ed_school[]": ["Columbia", "Stanford"],
        "ed_city[]":   ["New York", "Palo Alto"],
    }
    out = pack_repeating(form, "ed", ["school", "city"])
    assert out == [
        {"school": "Columbia", "city": "New York"},
        {"school": "Stanford", "city": "Palo Alto"},
    ]


def test_pack_repeating_skips_empty_rows():
    form = {"ed_school[]": ["Columbia", ""], "ed_city[]": ["NY", ""]}
    out = pack_repeating(form, "ed", ["school", "city"])
    assert len(out) == 1
    assert out[0]["school"] == "Columbia"


# --------- the real meat: run-preserving replacement ---------

def test_replace_in_single_run():
    doc, p = _new_doc_with_paragraph([("Hello [Name]!", False)])
    replace_in_runs_preserve(p, {"[Name]": "Alex"})
    assert p.text == "Hello Alex!"


def test_replace_preserves_other_runs_formatting():
    """Replacement must not touch runs that didn't contain the placeholder."""
    doc, p = _new_doc_with_paragraph([
        ("Name: ", False),
        ("[Name]", False),
        (" — ", False),
        ("Senior Analyst", True),   # bold — must stay bold
    ])
    replace_in_runs_preserve(p, {"[Name]": "Alex Chen"})
    assert p.text == "Name: Alex Chen — Senior Analyst"
    # The bold run must still be bold.
    bold_runs = [r for r in p.runs if r.bold]
    assert any("Senior Analyst" in r.text for r in bold_runs)


def test_replace_placeholder_split_across_runs():
    """Word frequently splits a placeholder into several runs."""
    doc, p = _new_doc_with_paragraph([
        ("[Com", False),
        ("pany ", False),
        ("Name]", False),
    ])
    replace_in_runs_preserve(p, {"[Company Name]": "Morgan Stanley"})
    assert p.text == "Morgan Stanley"


def test_longest_match_wins_when_keys_overlap():
    """`[Position Title], [Group Name]` must beat `[Position Title]`."""
    doc, p = _new_doc_with_paragraph([("[Position Title], [Group Name]", False)])
    replace_in_runs_preserve(p, {
        "[Position Title]": "WRONG",
        "[Position Title], [Group Name]": "Summer Analyst, M&A",
    })
    assert p.text == "Summer Analyst, M&A"


def test_none_values_become_empty_strings():
    doc, p = _new_doc_with_paragraph([("A [X] B", False)])
    replace_in_runs_preserve(p, {"[X]": None})
    assert p.text == "A  B"


def test_no_matches_is_a_noop():
    doc, p = _new_doc_with_paragraph([("nothing to replace", True)])
    before = p.text
    replace_in_runs_preserve(p, {"[X]": "Y"})
    assert p.text == before


def test_empty_paragraph_is_safe():
    doc = Document()
    p = doc.add_paragraph()
    replace_in_runs_preserve(p, {"[X]": "Y"})
    assert p.text == ""


def test_roundtrip_saves_as_valid_docx():
    """Rendered document should round-trip through save/load cleanly."""
    doc, p = _new_doc_with_paragraph([("Hello [Name]", False)])
    replace_in_runs_preserve(p, {"[Name]": "Alex"})

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    reloaded = Document(buf)
    assert reloaded.paragraphs[0].text == "Hello Alex"
