import io
import re

import pytest
from docx import Document

from app import app
from resume_builder.services import llm
from resume_builder.templates_registry import list_templates


@pytest.fixture()
def client():
    app.config["TESTING"] = True
    return app.test_client()


def _token(client):
    html = client.get("/build").get_data(as_text=True)
    return re.search(r'id="csrfToken" value="([^"]+)"', html).group(1)


def _cv(**personal):
    return {
        "style": {"template": "clean", "accent": "#2F5D8A", "font": "inter", "size": "m", "page": "A4"},
        "personal": {"firstName": "Alex", "lastName": "Chen", "email": "a@b.co", **personal},
        "sections": [
            {"type": "summary", "title": "Profile", "items": [{"text": "Analyst who likes numbers."}]},
            {"type": "experience", "title": "Experience", "items": [
                {"role": "Analyst", "org": "Acme", "start": "2023", "end": "Present", "description": "Built models\nCut costs 10%"},
            ]},
            {"type": "skills", "title": "Skills", "items": [{"name": "Excel", "level": "Expert"}, {"name": "Python"}]},
            {"type": "interests", "title": "Hidden bit", "hidden": True, "items": [{"name": "Chess"}]},
        ],
        "letter": {"company": "Acme", "role": "Analyst", "body": "First para.\n\nSecond para."},
    }


def test_has_many_templates_across_categories():
    tpls = list_templates()
    assert len(tpls) >= 20
    assert len({t.category for t in tpls}) >= 5


def test_favicon(client):
    r = client.get("/favicon.ico")
    assert r.status_code == 200 and r.mimetype == "image/png"
    html = client.get("/").get_data(as_text=True)
    assert "img/favicon.svg" in html and "apple-touch-icon" in html


def test_template_ids_unique_and_styled():
    import pathlib
    css = (pathlib.Path(__file__).resolve().parents[1] / "static/cv/cv.css").read_text()
    ids = [t.id for t in list_templates()]
    assert len(ids) == len(set(ids)) >= 45
    missing = [i for i in ids if f".cv--{i}" not in css]
    assert not missing, missing


def test_original_formats_are_back():
    by_id = {t.id: t for t in list_templates()}
    for tid in ("finance", "generic", "techline"):
        t = by_id[tid]
        assert t.category == "Original" and t.org_first and t.lines


def test_docx_export_in_times_for_finance(client):
    cv = _cv()
    cv["style"].update(template="finance", font="times")
    r = client.post("/api/export/docx", json={"cv": cv, "doc": "cv"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 200
    assert Document(io.BytesIO(r.data)).styles["Normal"].font.name == "Times New Roman"


def test_home_shows_featured_live_previews(client):
    html = client.get("/").get_data(as_text=True)
    assert 'class="cv-mini"' in html
    assert 'href="/templates"' in html
    assert "window.CV_META" in html


def test_gallery_lists_every_template(client):
    html = client.get("/templates").get_data(as_text=True)
    for t in list_templates():
        assert f'href="/build?template={t.id}"' in html


def test_build_passes_requested_template(client):
    html = client.get("/build?template=creative").get_data(as_text=True)
    assert 'requestedTemplate: "creative"' in html
    assert 'id="paper"' in html


def test_build_ignores_unknown_template(client):
    html = client.get("/build?template=nope").get_data(as_text=True)
    assert "requestedTemplate: null" in html


def test_docx_export(client):
    r = client.post("/api/export/docx", json={"cv": _cv(), "doc": "cv"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 200
    assert "Alex_Chen_CV.docx" in r.headers["Content-Disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    assert "Alex Chen" in text and "Cut costs 10%" in text and "Excel (Expert), Python" in text
    assert "Chess" not in text


def test_docx_cover_letter(client):
    r = client.post("/api/export/docx", json={"cv": _cv(), "doc": "letter"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 200
    assert "Alex_Chen_Cover_Letter.docx" in r.headers["Content-Disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    assert "Second para." in text and "Re: Analyst" in text


def test_docx_export_survives_odd_input(client):
    cv = {"style": {"accent": "javascript:alert(1)", "font": "nope", "page": "Huge"}, "personal": {"firstName": "../../x"},
          "sections": [None, {"type": "experience", "items": ["bad", {"role": 5}]}]}
    r = client.post("/api/export/docx", json={"cv": cv}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 200
    assert "/" not in r.headers["Content-Disposition"].split("filename=")[1]


def test_docx_export_needs_csrf(client):
    client.get("/build")
    r = client.post("/api/export/docx", json={"cv": _cv()})
    assert r.status_code in (400, 403)


def test_docx_export_rejects_bad_payload(client):
    r = client.post("/api/export/docx", json={"cv": "nope"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 400


def test_rewrite_rejects_unknown_mode(client, monkeypatch):
    monkeypatch.setenv("LLM_API_KEY", "test")
    r = client.post("/api/llm/rewrite-bullet", json={"text": "Did things", "mode": "shout"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 400


def test_write_summary_needs_some_input(client, monkeypatch):
    monkeypatch.setenv("LLM_API_KEY", "test")
    r = client.post("/api/llm/write-summary", json={}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 400


def test_ai_flag_follows_api_key(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    html = client.get("/build").get_data(as_text=True)
    assert "serverAi: false" in html and '"id": "anthropic"' in html
    monkeypatch.setenv("LLM_API_KEY", "test")
    assert "serverAi: true" in client.get("/build").get_data(as_text=True)


def _byok(client, provider="anthropic", key="sk-ant-test-key-123", model=""):
    return {"X-CSRF-Token": _token(client), "X-AI-Provider": provider, "X-AI-Key": key, "X-AI-Model": model}


def test_ai_without_any_key_asks_for_one(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    r = client.post("/api/llm/write-bullets", json={"role": "Analyst"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 401 and r.get_json()["needsKey"] is True


def test_ai_uses_the_visitors_own_key(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    seen = {}

    def fake_chat(cfg, system, user, **kw):
        seen["cfg"], seen["user"] = cfg, user
        return "Built models\nCut costs"

    monkeypatch.setattr(llm, "_chat", fake_chat)
    r = client.post("/api/llm/write-bullets", json={"role": "Analyst", "notes": "models", "job": "Bank analyst role"},
                    headers=_byok(client))
    assert r.status_code == 200 and r.get_json()["result"].startswith("Built")
    assert seen["cfg"].kind == "anthropic" and seen["cfg"].api_key == "sk-ant-test-key-123"
    assert seen["cfg"].model == "claude-opus-5-5"
    assert "Bank analyst role" in seen["user"]


def test_ai_model_override_and_openai_provider(client, monkeypatch):
    seen = {}
    monkeypatch.setattr(llm, "_chat", lambda cfg, *a, **k: seen.setdefault("cfg", cfg) and "ok")
    r = client.post("/api/llm/test", json={}, headers=_byok(client, "openrouter", "sk-or-abcdefgh", "meta/llama-3"))
    assert r.status_code == 200
    assert seen["cfg"].base_url == "https://openrouter.ai/api/v1" and seen["cfg"].model == "meta/llama-3"


@pytest.mark.parametrize("provider,key,model", [("evil", "sk-abcdefgh", ""), ("openai", "short", ""), ("openai", "sk-abcdefgh", "bad model!")])
def test_ai_rejects_bad_settings(client, provider, key, model):
    r = client.post("/api/llm/test", json={}, headers=_byok(client, provider, key, model))
    assert r.status_code == 400


def test_suggest_skills_skips_ones_already_listed(client, monkeypatch):
    monkeypatch.setattr(llm, "_chat", lambda *a, **k: "1. Excel, Python\n- Financial modelling, Bloomberg")
    r = client.post("/api/llm/suggest-skills", json={"headline": "Analyst", "have": "Excel, python"}, headers=_byok(client))
    assert r.get_json()["result"] == ["Financial modelling", "Bloomberg"]


def test_ai_errors_are_friendly(client, monkeypatch):
    def boom(*a, **k):
        raise llm.LLMError("Claude rejected that API key. Check it in AI settings.")
    monkeypatch.setattr(llm, "_chat", boom)
    r = client.post("/api/llm/test", json={}, headers=_byok(client))
    assert r.status_code == 502 and "rejected" in r.get_json()["error"]


SAMPLE_CV_TEXT = """Jordan Patel
Software Engineer
jordan.patel@example.com | +44 7700 900456 | Leeds, UK | linkedin.com/in/jordanpatel

Profile
Backend engineer who enjoys making slow systems fast.

Experience
Software Engineer
Sky Betting & Gaming | Leeds | Sep 2022 – Present
• Cut API latency by 40% by adding a Redis cache
• Led the move from cron jobs to a queue

Education
BSc Computer Science
University of Leeds | 2019 – 2022
First Class Honours

Skills
Python (Expert), Go (Intermediate), PostgreSQL

Languages
English (Native), Gujarati (Fluent)
"""


def _upload(client, name, data, **headers):
    return client.post("/api/import", data={"file": (io.BytesIO(data), name)},
                       headers={"X-CSRF-Token": _token(client), **headers}, content_type="multipart/form-data")


def _section(cv, kind):
    return next(s for s in cv["sections"] if s["type"] == kind)


def test_import_text_with_basic_reader(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    r = _upload(client, "cv.txt", SAMPLE_CV_TEXT.encode())
    assert r.status_code == 200
    body = r.get_json()
    cv = body["cv"]
    assert body["usedAi"] is False
    p = cv["personal"]
    assert (p["firstName"], p["lastName"], p["headline"]) == ("Jordan", "Patel", "Software Engineer")
    assert p["email"] == "jordan.patel@example.com" and "linkedin.com/in/jordanpatel" in p["linkedin"]
    job = _section(cv, "experience")["items"][0]
    assert job["role"] == "Software Engineer" and job["org"] == "Sky Betting & Gaming"
    assert (job["start"], job["end"]) == ("Sep 2022", "Present")
    assert job["description"].splitlines() == ["Cut API latency by 40% by adding a Redis cache", "Led the move from cron jobs to a queue"]
    edu = _section(cv, "education")["items"][0]
    assert edu["degree"] == "BSc Computer Science" and edu["org"] == "University of Leeds" and "First Class" in edu["grade"]
    skills = {i["name"]: i.get("level", "") for i in _section(cv, "skills")["items"]}
    assert skills == {"Python": "Expert", "Go": "Intermediate", "PostgreSQL": ""}
    langs = {i["name"]: i.get("level", "") for i in _section(cv, "languages")["items"]}
    assert langs == {"English": "Native", "Gujarati": "Fluent"}


def test_import_word_document(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    doc = Document()
    for line in SAMPLE_CV_TEXT.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    r = _upload(client, "My CV.docx", buf.getvalue())
    assert r.status_code == 200
    cv = r.get_json()["cv"]
    assert cv["personal"]["lastName"] == "Patel"
    assert _section(cv, "experience")["items"][0]["org"] == "Sky Betting & Gaming"


@pytest.mark.parametrize("name,data", [("cv.exe", b"MZ" * 40), ("cv.pdf", b"not really a pdf"), ("cv.txt", b"   ")])
def test_import_rejects_unreadable_files(client, name, data):
    r = _upload(client, name, data)
    assert r.status_code == 400 and r.get_json()["error"]


def test_import_needs_csrf_and_a_file(client):
    r = client.post("/api/import", data={"file": (io.BytesIO(b"x" * 50), "cv.txt")}, content_type="multipart/form-data")
    assert r.status_code in (400, 403)
    r = client.post("/api/import", data={}, headers={"X-CSRF-Token": _token(client)}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_import_with_ai_keeps_only_known_fields(client, monkeypatch):
    reply = """```json
    {"personal": {"firstName": "Jordan", "lastName": "Patel", "password": "x"},
     "sections": [
       {"type": "experience", "title": "Experience", "items": [{"role": "Engineer", "org": "Sky", "salary": "lots"}]},
       {"type": "skills", "title": "Skills", "items": [{"name": "Python", "level": "Wizard"}]},
       {"type": "malware", "title": "Nope", "items": [{"name": "x"}]}
     ]}
    ```"""
    monkeypatch.setattr(llm, "_chat", lambda *a, **k: reply)
    r = _upload(client, "cv.txt", SAMPLE_CV_TEXT.encode(), **{k: v for k, v in _byok(client).items() if k != "X-CSRF-Token"})
    body = r.get_json()
    assert r.status_code == 200 and body["usedAi"] is True
    cv = body["cv"]
    assert "password" not in cv["personal"]
    # Unknown section types are kept as a plain custom section, under their own heading.
    assert [(s["type"], s["title"]) for s in cv["sections"]] == [("experience", "Experience"), ("skills", "Skills"), ("custom", "Nope")]
    assert "salary" not in cv["sections"][0]["items"][0]
    assert cv["sections"][1]["items"][0]["level"] == ""


def test_import_falls_back_when_ai_fails(client, monkeypatch):
    def boom(*a, **k):
        raise llm.LLMError("Claude rejected that API key. Check it in AI settings.")
    monkeypatch.setattr(llm, "_chat", boom)
    r = _upload(client, "cv.txt", SAMPLE_CV_TEXT.encode(), **{k: v for k, v in _byok(client).items() if k != "X-CSRF-Token"})
    body = r.get_json()
    assert r.status_code == 200 and body["usedAi"] is False and "basic reader" in body["note"]
    assert body["cv"]["personal"]["firstName"] == "Jordan"


def test_translate_returns_only_the_keys_it_was_given(client, monkeypatch):
    seen = {}

    def fake_chat(cfg, system, user, **kw):
        seen["system"] = system
        return '{"p.headline": "Ingeniera de software", "s.exp": "Experiencia", "extra": "no"}'

    monkeypatch.setattr(llm, "_chat", fake_chat)
    r = client.post("/api/llm/translate", json={"strings": {"p.headline": "Software engineer", "s.exp": "Experience"},
                                                "language": "Spanish"}, headers=_byok(client))
    assert r.status_code == 200
    assert r.get_json()["result"] == {"p.headline": "Ingeniera de software", "s.exp": "Experiencia"}
    assert "Spanish" in seen["system"]


@pytest.mark.parametrize("payload", [{}, {"strings": "hi", "language": "French"}, {"strings": {"a": "b"}, "language": ""}])
def test_translate_rejects_bad_payload(client, payload):
    r = client.post("/api/llm/translate", json=payload, headers=_byok(client))
    assert r.status_code == 400


def test_templates_say_which_are_ats_friendly(client):
    tpls = {t.id: t for t in list_templates()}
    assert tpls["clean"].ats and tpls["finance"].ats
    assert not tpls["vertical"].ats
    assert not any(t.ats for t in tpls.values() if t.photo or t.layout != "single")
    assert '"ats": true' in client.get("/build").get_data(as_text=True)


# --- CV import: hostile files stay cheap, odd parses never 500 ---------------

def _timed_upload(client, name, data, **headers):
    import time
    t = time.perf_counter()
    r = _upload(client, name, data, **headers)
    assert time.perf_counter() - t < 2, f"import of {name} took too long"
    return r


def _docx_with_body(body_xml):
    import zipfile
    buf = io.BytesIO()
    Document().save(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                data = (b'<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        b"<w:body>" + body_xml + b"</w:body></w:document>")
            zout.writestr(info, data)
    return out.getvalue()


def _pdf_with_content(content):
    import zlib
    stream = zlib.compress(content)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    pdf, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n" % (i + 1) + o + b"\nendobj\n"
    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return pdf + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)


@pytest.mark.parametrize("text", [
    "John Smith\njohn@example.com\nSkills\nPython Expert Java Expert a" + "　" * 3000 + "b",
    "John Smith\njohn@example.com\nSkills\nPython Expert Java Advanced\n" + "\n".join(["more stuff here"] * 1200),
])
def test_import_skill_runs_stay_fast(client, monkeypatch, text):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    r = _timed_upload(client, "cv.txt", text.encode())
    assert r.status_code == 200 and r.get_json()["cv"]["personal"]["lastName"] == "Smith"


def test_import_skips_huge_pdf_pages(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    pdf = _pdf_with_content(b"BT /F1 12 Tf 72 700 Td (hello world) Tj ET\n" * 100000)  # 4 MB of drawing on one page
    r = _timed_upload(client, "cv.pdf", pdf)
    assert r.status_code == 400 and r.get_json()["error"]


@pytest.mark.parametrize("body", [
    b"<w:p/>" * 1_700_000,  # tiny zip, 10 MB of empty paragraphs
    b"<w:tbl><w:tr><w:tc><w:tcPr><w:vMerge w:val='restart'/></w:tcPr><w:p><w:r><w:t>Jane Smith jane@example.com</w:t></w:r></w:p>"
    b"</w:tc></w:tr>" + b"<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>" * 1000 + b"</w:tbl>",
])
def test_import_hostile_word_files_are_cheap(client, monkeypatch, body):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    r = _timed_upload(client, "cv.docx", _docx_with_body(body))
    assert r.status_code in (200, 400) and r.is_json


def test_import_unexpected_reader_error_is_a_friendly_400(client, monkeypatch):
    from resume_builder.services import cv_import

    def boom(*a, **k):
        raise RecursionError("deep")
    monkeypatch.setattr(cv_import, "extract_text", boom)
    r = _upload(client, "cv.txt", SAMPLE_CV_TEXT.encode())
    assert r.status_code == 400 and r.get_json()["error"]


@pytest.mark.parametrize("reply", ['{"sections": [{"type": ["experience"]}]}', '{"sections": 1}',
                                   '{"personal": {"firstName": "Jo"}, "sections": [null, "x"]}'])
def test_import_odd_ai_replies_fall_back(client, monkeypatch, reply):
    monkeypatch.setattr(llm, "_chat", lambda *a, **k: reply)
    r = _upload(client, "cv.txt", SAMPLE_CV_TEXT.encode(), **{k: v for k, v in _byok(client).items() if k != "X-CSRF-Token"})
    body = r.get_json()
    assert r.status_code == 200 and body["usedAi"] is False and body["cv"]["personal"]["lastName"] == "Patel"


def test_import_word_keeps_reading_order(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Daniel Kim"
    doc.add_paragraph("daniel.kim@example.com")
    doc.add_paragraph("Experience")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "Product Manager | Spotify", "Jan 2021 - Present"
    doc.add_paragraph("Shipped Playlist Search To 200M Users", style="List Bullet")
    doc.add_paragraph("Skills")
    doc.add_paragraph("SQL, Figma")
    buf = io.BytesIO()
    doc.save(buf)
    cv = _upload(client, "cv.docx", buf.getvalue()).get_json()["cv"]
    assert (cv["personal"]["firstName"], cv["personal"]["lastName"]) == ("Daniel", "Kim")
    job = _section(cv, "experience")["items"][0]
    assert (job["role"], job["org"], job["start"], job["end"]) == ("Product Manager", "Spotify", "Jan 2021", "Present")
    assert job["description"] == "Shipped Playlist Search To 200M Users"
    assert [i["name"] for i in _section(cv, "skills")["items"]] == ["SQL", "Figma"]


def test_import_basic_reader_edge_cases(client, monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    text = """﻿CURRICULUM VITAE
JEAN-PIERRE DUPONT
(212) 555-0199 | jp@example.com
Experience
Senior Software Engineer | Stripe | San Francisco, CA | Mar 2020 - Present
 Promoted to team lead in 2021
 Shipped v2.0 in March 2020
Analyst
Barclays, London
2018 - 2019
• Built credit models
Research Assistant
UCL
Fall 2017 - Spring 2018
• Ran lab studies
Education
BSc Mathematics
University of Leeds
2015 - 2018
GPA: 3.9/4.0
Technical Skills
Languages: Java, Python, SQL (Postgres)
Frameworks: React
Publications
Dupont J. A study of things. Nature, 2021.
"""
    cv = _upload(client, "cv.txt", text.encode()).get_json()["cv"]
    p = cv["personal"]
    assert (p["firstName"], p["lastName"], p["phone"]) == ("Jean-Pierre", "Dupont", "(212) 555-0199")
    jobs = _section(cv, "experience")["items"]
    assert [(j["role"], j["org"], j["location"], j["start"], j["end"]) for j in jobs] == [
        ("Senior Software Engineer", "Stripe", "San Francisco, CA", "Mar 2020", "Present"),
        ("Analyst", "Barclays", "London", "2018", "2019"),
        ("Research Assistant", "UCL", "", "Fall 2017", "Spring 2018"),
    ]
    assert jobs[0]["description"].splitlines() == ["Promoted to team lead in 2021", "Shipped v2.0 in March 2020"]
    edu = _section(cv, "education")["items"][0]
    assert (edu["grade"], edu["description"]) == ("GPA: 3.9/4.0", "")
    assert [i["name"] for i in _section(cv, "skills")["items"]] == ["Java", "Python", "SQL (Postgres)", "React"]
    assert not any(s["type"] == "languages" for s in cv["sections"])
    assert _section(cv, "custom")["title"] == "Publications"


@pytest.mark.parametrize("head,role,org", [
    ("Software Engineer, Google", "Software Engineer", "Google"),
    ("Data Analyst at Monzo", "Data Analyst", "Monzo"),
    ("Summer Analyst, Equity Research", "Summer Analyst, Equity Research", ""),
])
def test_import_splits_role_and_company_on_one_line(head, role, org):
    from resume_builder.services import cv_import
    text = f"Jane Doe\njane@example.com\n\nExperience\n{head}\nJan 2020 - Present\n• Built a reporting tool"
    cv = cv_import.normalize(cv_import.heuristic_parse(text))
    job = _section(cv, "experience")["items"][0]
    assert (job["role"], job["org"]) == (role, org)


def test_docx_letter_uses_translated_greeting(client):
    cv = _cv()
    cv["letter"].update({"greet": "Estimado/a", "anyone": "responsable de selección", "reLabel": "Asunto:"})
    r = client.post("/api/export/docx", json={"cv": cv, "doc": "letter"}, headers={"X-CSRF-Token": _token(client)})
    assert r.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(r.data)).paragraphs)
    assert "Asunto: Analyst" in text and "Estimado/a responsable de selección," in text and "Dear" not in text
