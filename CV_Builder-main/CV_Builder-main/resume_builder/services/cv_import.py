"""Turn an uploaded CV (PDF, Word or plain text) into the editor's data shape.

``extract_text`` pulls the words out of the file. ``heuristic_parse`` makes a
best guess at the structure without any AI: contact details by pattern,
sections by their usual headings, entries by dates and bullet points.
``normalize`` cleans any parsed result (ours or the AI's) down to the fields
and section types the editor knows, so nothing unexpected reaches the page.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Dict, List, Optional

MAX_TEXT = 20000

SECTION_FIELDS = {
    "summary": ["text"],
    "experience": ["role", "org", "location", "start", "end", "description"],
    "education": ["degree", "org", "location", "start", "end", "grade", "description"],
    "projects": ["name", "link", "start", "end", "description"],
    "volunteering": ["role", "org", "location", "start", "end", "description"],
    "skills": ["name", "level"],
    "languages": ["name", "level"],
    "certifications": ["name", "org", "end"],
    "awards": ["name", "org", "end", "description"],
    "interests": ["name"],
    "custom": ["name", "org", "location", "start", "end", "description"],
}
TITLES = {
    "summary": "Profile", "experience": "Experience", "education": "Education", "projects": "Projects",
    "volunteering": "Volunteering", "skills": "Skills", "languages": "Languages",
    "certifications": "Certifications", "awards": "Awards", "interests": "Interests", "custom": "Other",
}
PERSONAL_FIELDS = ["firstName", "lastName", "headline", "email", "phone", "location", "website", "linkedin"]
SKILL_LEVELS = ["Beginner", "Intermediate", "Advanced", "Expert"]
LANGUAGE_LEVELS = ["Native", "Fluent", "Advanced", "Intermediate", "Basic"]


class ImportError_(ValueError):
    """The file couldn't be read."""


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        text = _pdf_text(data)
    elif name.endswith(".docx") or data[:2] == b"PK":
        text = _docx_text(data)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ImportError_("Upload a PDF, Word (.docx) or text file.")
    text = text.replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()
    if len(text) < 30:
        raise ImportError_("We couldn't find any text in that file. If it's a scanned image, try the Word version.")
    return text[:MAX_TEXT]


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:6])
    except Exception as e:  # pypdf raises many types for broken files
        raise ImportError_("That PDF couldn't be read.") from e


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as z:  # refuse files that unpack to something huge
            if sum(i.file_size for i in z.infolist()) > 60 * 1024 * 1024:
                raise ImportError_("That Word file is too large to read.")
        doc = Document(io.BytesIO(data))
    except ImportError_:
        raise
    except Exception as e:
        raise ImportError_("That Word file couldn't be read. Save it as .docx and try again.") from e
    lines: List[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:
                    cells.append(t)
            lines.append("  |  ".join(cells))
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Heuristic parsing (no AI)
# ---------------------------------------------------------------------------

EMAIL = re.compile(r"[\w.+\-]+@[\w\-]+(\.[\w\-]+)+")
PHONE = re.compile(r"(\+?\d[\d\s().\-]{7,}\d)")
LINKEDIN = re.compile(r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/[^\s|,]+", re.I)
WEBSITE = re.compile(r"(?:https?://)?(?:www\.)?[a-z0-9\-]+\.(?:com|io|dev|me|net|org|co\.uk|co|app|site|xyz)(?:/[^\s|,]*)?", re.I)
BULLET = re.compile(r"^\s*[•●▪◦‣∙·\-–—*»>]\s*")
MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE = rf"(?:{MONTH}\s*\d{{4}}|\d{{1,2}}/\d{{4}}|\d{{4}}/\d{{1,2}}|(?:19|20)\d{{2}})"
RANGE = re.compile(rf"({DATE})\s*(?:-|–|—|to|until)\s*({DATE}|present|current|now|today)|(?:expected\s+)?({DATE})", re.I)
LEVELS = ["Expert", "Advanced", "Intermediate", "Beginner", "Native", "Fluent", "Basic"]
LEVEL_RUN = re.compile(r"\s*(.+?)\s*[(\-–:]?\s*\b(" + "|".join(LEVELS) + r")\b\)?\s*(?:,|;|$|(?=\s))", re.I)
ROLE_WORDS = re.compile(
    r"\b(analyst|engineer|manager|intern|internship|president|director|assistant|developer|consultant|associate|officer|"
    r"lead|head|designer|specialist|coordinator|teacher|tutor|scientist|researcher|member|volunteer|founder|co-founder|"
    r"executive|representative|administrator|accountant|nurse|editor|writer|chair|captain|ambassador|advisor|adviser|"
    r"technician|programmer|architect|owner|partner|vp|ceo|cto|cfo|trainee|apprentice|clerk|cashier|server|barista|"
    r"supervisor|instructor|fellow|student|mentor|organiser|organizer|treasurer|secretary)\b", re.I)
ORG_WORDS = re.compile(r"\b(university|college|school|institute|academy|ltd|limited|inc|llc|plc|group|bank|society|club|"
                       r"foundation|company|corp|corporation|council|association|agency|partners|capital|labs?)\b", re.I)
DEGREE = re.compile(r"\b(B\.?Sc|B\.?A|BEng|MEng|M\.?Sc|M\.?A|MBA|Ph\.?D|LLB|LLM|Bachelor|Master|Doctor|Diploma|Certificate|"
                    r"A[- ]?Levels?|GCSE|IB|High School|Associate|Foundation Year|HND|BTEC)\b", re.I)
GRADE = re.compile(r"(GPA[:\s]*[\d.]+(?:\s*/\s*[\d.]+)?|First[- ]Class( Honours)?|Upper Second|2[:.]1|2[:.]2|Distinction|Merit|"
                   r"Summa Cum Laude|Magna Cum Laude|Cum Laude|\b[A-D]\*?[A-D]\*?[A-D]\*?\b)", re.I)
LOCATION = re.compile(r"^(remote|hybrid|[A-Z][a-zÀ-ÿ'.]+(?:[ \-][A-Z][a-zÀ-ÿ'.]+)?(?:,\s*[A-Z][A-Za-zÀ-ÿ .]+)?)$")

HEADINGS = [
    ("contact", r"contact( details| information| info)?|personal (details|information)"),
    ("summary", r"(professional |personal |career )?(summary|profile|about me|about|objective|introduction)"),
    ("experience", r"(work |professional |relevant |employment |career )?(experience|history|employment)( history)?|work (&|and) leadership experience|leadership experience"),
    ("education", r"education( (&|and) training)?|academic (background|qualifications)|qualifications"),
    ("projects", r"(personal |selected |key |academic )?projects"),
    ("volunteering", r"volunteer(ing)?( experience| work)?|community (work|involvement)"),
    ("skills", r"(technical |key |core |professional )?(skills|competencies|expertise|tools)( (&|and) tools)?|skills,? (activities|languages) (&|and) interests|skills (&|and) (interests|languages)|additional information"),
    ("languages", r"languages?"),
    ("certifications", r"certifications?|certificates?|licen[cs]es( (&|and) certifications)?|courses|training"),
    ("awards", r"awards?( (&|and) honou?rs)?|honou?rs( (&|and) awards)?|achievements|scholarships"),
    ("interests", r"interests|hobbies( (&|and) interests)?|activities"),
]
_HEAD_RE = [(t, re.compile(rf"^\s*(?:{p})\s*:?\s*$", re.I)) for t, p in HEADINGS]
LABELS = [  # "Label: a, b" lines inside a combined skills block
    ("languages", re.compile(r"^languages?\s*:\s*", re.I)),
    ("interests", re.compile(r"^(interests|hobbies|activities)\s*:\s*", re.I)),
    ("certifications", re.compile(r"^(certifications?|certificates?)( (&|and) training)?\s*:\s*", re.I)),
    ("skills", re.compile(r"^([A-Za-z &/]{0,25}skills|tools|software|programming( languages)?|technologies)\s*:\s*", re.I)),
]


def _heading_type(line: str) -> Optional[str]:
    if len(line) > 48:
        return None
    for t, rx in _HEAD_RE:
        if rx.match(line):
            return t
    return None


def _is_contact(line: str) -> bool:
    return bool(EMAIL.search(line) or LINKEDIN.search(line) or re.fullmatch(r"[+\d][\d\s().\-]{7,}", line.strip()))


def heuristic_parse(text: str) -> Dict[str, object]:
    # Real CV words and lines are short; capping them keeps the patterns below fast on odd files.
    text = re.sub(r"\S{120,}", lambda m: m.group(0)[:120], text[:MAX_TEXT])
    lines = [l.strip()[:1500] for l in text.split("\n")]
    lines = [l for l in lines if l]
    personal: Dict[str, str] = {}

    head_n = next((i for i, l in enumerate(lines) if _heading_type(l) and _heading_type(l) != "contact"), len(lines))
    top = "\n".join(lines[:max(head_n, 6)])
    every = "\n".join(lines)
    for src in (top, every):
        if "email" not in personal and (m := EMAIL.search(src)):
            personal["email"] = m.group(0)
        if "linkedin" not in personal and (m := LINKEDIN.search(src)):
            personal["linkedin"] = re.sub(r"^https?://(www\.)?", "", m.group(0))
        if "phone" not in personal:
            for cand in PHONE.findall(src):
                digits = re.sub(r"\D", "", cand)
                if 8 <= len(digits) <= 15 and not RANGE.fullmatch(cand.strip()):
                    personal["phone"] = re.sub(r"\s+", " ", cand.strip())
                    break
    for m in WEBSITE.finditer(top):
        url = m.group(0)
        if "linkedin" in url.lower() or url in personal.get("email", ""):
            continue
        personal["website"] = re.sub(r"^https?://(www\.)?", "", url)
        break

    header = lines[:head_n]
    name_idx = None
    for i, l in enumerate(header[:6]):
        if re.fullmatch(r"[A-Za-zÀ-ÿ'’.\- ]{3,50}", l) and 2 <= len(l.split()) <= 4:
            name_idx = i
            break
    if name_idx is not None:
        parts = header[name_idx].split()
        if all(p.isupper() for p in parts):
            parts = [p.capitalize() for p in parts]
        personal["firstName"] = " ".join(parts[:-1])
        personal["lastName"] = parts[-1]
        nxt = header[name_idx + 1] if name_idx + 1 < len(header) else ""
        if nxt and len(nxt) <= 70 and not _is_contact(nxt) and not PHONE.search(nxt) and "|" not in nxt:
            personal["headline"] = nxt
    for l in header:
        for seg in re.split(r"\s*[|•·]\s*", l):
            seg = seg.strip()
            if seg and not _is_contact(seg) and LOCATION.match(seg) and "," in seg and seg != personal.get("headline"):
                personal["location"] = seg
                break
        if "location" in personal:
            break

    blocks: List[Dict[str, object]] = []
    cur: Optional[Dict[str, object]] = None
    for l in lines[head_n:]:
        t = _heading_type(l)
        if t:
            title = l.strip(" :")
            cur = {"type": t, "title": title.title() if title.isupper() else title, "lines": []}
            blocks.append(cur)
        elif cur is not None and not (cur["type"] != "summary" and _is_contact(l)):
            cur["lines"].append(l)

    sections: List[Dict[str, object]] = []
    for b in blocks:
        if b["type"] == "contact":
            continue
        if b["type"] == "skills" and any(rx.match(l) for l in b["lines"] for _, rx in LABELS):
            sections.extend(_labelled_lists(b["lines"]))
            continue
        items = _parse_section(b["type"], b["lines"])
        if items:
            sections.append({"type": b["type"], "title": b["title"], "items": items})
    if not sections and len(lines) > (name_idx or 0) + 2:
        body = [l for l in lines[(name_idx or 0) + 1:] if not _is_contact(l)]
        sections.append({"type": "summary", "title": "Profile", "items": [{"text": " ".join(body)[:1500]}]})
    return {"personal": personal, "sections": sections}


def _join_wrapped(lines: List[str]) -> List[str]:
    """Undo PDF line wrapping inside lists: "Bloomberg" + "(Advanced), …"."""
    out: List[str] = []
    for l in lines:
        if out and (l[:1] in "(,&" or l[:1].islower() or out[-1].rstrip().endswith((",", "&", "/"))) and not BULLET.match(l):
            out[-1] = out[-1].rstrip() + " " + l
        else:
            out.append(l)
    return out


def _list_items(lines: List[str], kind: str) -> List[Dict[str, str]]:
    names: List[Dict[str, str]] = []
    for l in _join_wrapped(lines):
        l = BULLET.sub("", l)
        for _, rx in LABELS:
            l = rx.sub("", l)
        l = re.sub(r"^[A-Za-z &/]{2,30}:\s*", "", l)
        if kind != "interests" and len(re.findall(r"\b(" + "|".join(LEVELS) + r")\b", l, re.I)) >= 2 and not re.search(r"[,;|•]", l):
            parts = [(m.group(1), m.group(2)) for m in LEVEL_RUN.finditer(l)]
        else:
            parts = []
            for p in re.split(r"\s*[,;•|·]\s*", l):
                m = re.match(r"^(.*?)[\s(\-–:]+(" + "|".join(LEVELS) + r")\)?$", p.strip(), re.I)
                parts.append((m.group(1), m.group(2)) if m else (p, ""))
        for name, level in parts:
            name = name.strip(" .()")
            level = level.capitalize()
            if name and len(name) <= 60 and all(name.lower() != n["name"].lower() for n in names):
                item = {"name": name}
                if kind != "interests":
                    allowed = LANGUAGE_LEVELS if kind == "languages" else SKILL_LEVELS
                    item["level"] = level if level in allowed else ""
                names.append(item)
    return names[:40]


def _labelled_lists(lines: List[str]) -> List[Dict[str, object]]:
    groups: Dict[str, List[str]] = {}
    order: List[str] = []
    kind = "skills"
    for l in _join_wrapped(lines):
        for k, rx in LABELS:
            if rx.match(l):
                kind = k
                break
        if kind not in groups:
            groups[kind] = []
            order.append(kind)
        groups[kind].append(l)
    out = []
    for k in order:
        items = _list_items(groups[k], k) if k != "certifications" else [{"name": n["name"], "end": ""} for n in _list_items(groups[k], "interests")]
        if items:
            out.append({"type": k, "title": TITLES[k], "items": items})
    return out


def _dates(line: str):
    m = RANGE.search(line)
    if not m:
        return line, "", ""
    if m.group(1):
        start, end = m.group(1), m.group(2)
        end = "Present" if end.lower() in ("present", "current", "now", "today") else end
    else:
        start, end = "", m.group(3)
    rest = (line[:m.start()] + line[m.end():]).strip(" ,|–—-()·")
    return rest, start.strip(), end.strip()


def _is_heady(line: str) -> bool:
    """Short title-like line (employer, role, school, place) rather than a sentence."""
    if BULLET.match(line) or ":" in line or line.endswith("."):
        return False
    words = re.findall(r"[A-Za-zÀ-ÿ][\w'’&.\-]*", line)
    if not words or len(words) > 8:
        return False
    small = {"of", "and", "in", "the", "for", "at", "&", "de", "a", "an", "on", "to"}
    caps = [w for w in words if w.lower() not in small]
    return bool(caps) and sum(w[0].isupper() for w in caps) / len(caps) >= 0.6


def _parse_section(kind: str, lines: List[str]) -> List[Dict[str, str]]:
    if not lines:
        return []
    if kind == "summary":
        return [{"text": " ".join(BULLET.sub("", l) for l in lines)[:1500]}]
    if kind in ("skills", "interests", "languages"):
        return _list_items(lines, kind)
    if kind in ("certifications", "awards"):
        items = []
        for l in _join_wrapped(lines):
            rest, _start, end = _dates(BULLET.sub("", l))
            if rest:
                items.append({"name": rest[:120], "end": end})
        return items[:20]

    # Entries: a run of title-like lines (with the dates somewhere) and then
    # the description lines. A title-like line after a description starts the
    # next entry.
    entries: List[Dict[str, object]] = []
    cur: Optional[Dict[str, object]] = None
    for l in lines:
        rest, start, end = _dates(l)
        dated = bool(start or end)
        heady = dated or _is_heady(l)
        if heady and (cur is None or cur["desc"] or (dated and (cur["start"] or cur["end"]))):
            cur = {"head": [], "desc": [], "start": "", "end": ""}
            entries.append(cur)
        elif cur is None:
            cur = {"head": [], "desc": [], "start": "", "end": ""}
            entries.append(cur)
        if heady:
            if dated:
                cur["start"], cur["end"] = cur["start"] or start, cur["end"] or end
            for p in re.split(r"\s+[|–—]\s+|\s{3,}|\s+·\s+", rest):
                if p.strip():
                    cur["head"].append(p.strip(" ,"))
        elif cur["desc"] and l[:1].islower() and not BULLET.match(l):
            cur["desc"][-1] += " " + l
        else:
            cur["desc"].append(BULLET.sub("", l))

    items = []
    for e in entries[:15]:
        head = [h for h in e["head"] if h]
        desc = "\n".join(e["desc"])[:1500]
        places = [h for h in head if LOCATION.match(h) and not ROLE_WORDS.search(h) and not ORG_WORDS.search(h)
                  and (len(head) > 2 or "," in h or h.lower() in ("remote", "hybrid"))]
        places.sort(key=lambda h: (len(h.split(",")[0].split()), "," not in h))
        loc = places[0] if places and (len(places[0].split(",")[0].split()) == 1 or "," in places[0]) else ""
        rest = [h for h in head if h != loc]
        if kind == "education":
            grade = ""
            for i, h in enumerate(rest):
                gm = GRADE.search(h)
                if gm and (len(h) - len(gm.group(0)) < 4 or "," in h or "·" in h):
                    grade = gm.group(0)
                    rest[i] = h.replace(gm.group(0), "").strip(" ,·-")
            if not grade and (gm := GRADE.search(desc)) and gm.group(0).lower().startswith(("gpa", "first", "upper", "2")):
                grade = gm.group(0)
            rest = [h for h in rest if h]
            degree = next((h for h in rest if DEGREE.search(h)), rest[0] if rest else "")
            org = next((h for h in rest if h != degree), "")
            items.append({"degree": degree, "org": org, "location": loc, "start": e["start"], "end": e["end"], "grade": grade, "description": desc})
        elif kind == "projects":
            name = rest[0] if rest else ""
            link = next((h for h in rest[1:] if re.search(r"\.\w{2,}/?", h) and " " not in h), "")
            items.append({"name": name, "link": link, "start": e["start"], "end": e["end"], "description": desc})
        else:
            role = next((h for h in rest if ROLE_WORDS.search(h) and not ORG_WORDS.search(h)), "")
            if not role and rest:
                role = rest[0] if len(rest) == 1 or not ORG_WORDS.search(rest[0]) else rest[1]
            others = [h for h in rest if h != role]
            org = others[0] if others else ""
            extra = others[1:]
            if extra:
                desc = "\n".join(extra + ([desc] if desc else []))
            items.append({"role": role, "org": org, "location": loc, "start": e["start"], "end": e["end"], "description": desc})
    return [i for i in items if any(str(v).strip() for v in i.values())]


# ---------------------------------------------------------------------------
# Normalising (shared by the heuristic and AI paths)
# ---------------------------------------------------------------------------

def _s(v, limit: int = 200) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        v = "\n".join(str(x) for x in v)
    return str(v).strip()[:limit]


def normalize(data: object) -> Dict[str, object]:
    if not isinstance(data, dict):
        raise ImportError_("The CV couldn't be understood.")
    p_in = data.get("personal") if isinstance(data.get("personal"), dict) else {}
    personal = {k: _s(p_in.get(k), 120) for k in PERSONAL_FIELDS}
    sections = []
    for sec in data.get("sections") or []:
        if not isinstance(sec, dict):
            continue
        kind = sec.get("type") if sec.get("type") in SECTION_FIELDS else "custom"
        fields = SECTION_FIELDS[kind]
        items = []
        raw_items = sec.get("items")
        for it in (raw_items if isinstance(raw_items, list) else [])[:40]:
            if isinstance(it, str):
                it = {fields[0]: it}
            if not isinstance(it, dict):
                continue
            clean = {k: _s(it.get(k), 2000 if k in ("description", "text") else 200) for k in fields}
            if kind == "skills" and clean.get("level") not in SKILL_LEVELS:
                clean["level"] = ""
            if kind == "languages" and clean.get("level") not in LANGUAGE_LEVELS:
                clean["level"] = ""
            if any(clean.values()):
                items.append(clean)
        if items:
            sections.append({"type": kind, "title": _s(sec.get("title"), 60) or TITLES[kind], "items": items})
    if kind_count(sections, "summary") > 1:
        first = next(s for s in sections if s["type"] == "summary")
        sections = [s for s in sections if s["type"] != "summary" or s is first]
    return {"personal": personal, "sections": sections[:20]}


def kind_count(sections: List[Dict[str, object]], kind: str) -> int:
    return sum(1 for s in sections if s["type"] == kind)
