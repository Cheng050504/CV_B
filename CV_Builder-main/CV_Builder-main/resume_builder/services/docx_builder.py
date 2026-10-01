"""Build a Word (.docx) CV or cover letter from the editor's CV data.

The browser draws each template exactly (and prints the PDF); Word cannot
reproduce every layout, so this produces one clean, ATS-friendly single
column that keeps the user's section order, accent colour and font family.
"""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..templates_registry import FONTS

_HEX = re.compile(r"^#?([0-9a-fA-F]{6})$")
PAGE_SIZES = {"A4": (Mm(210), Mm(297)), "Letter": (Mm(215.9), Mm(279.4))}
BASE_PT = {"s": 9.5, "m": 10.5, "l": 11.5}

# Mirrors SECTION_TYPES in static/cv/engine.js: which keys form the title / subtitle.
_ENTRY_KEYS = {
    "experience": ("role", "org"),
    "education": ("degree", "org"),
    "projects": ("name", "link"),
    "volunteering": ("role", "org"),
    "certifications": ("name", "org"),
    "awards": ("name", "org"),
    "custom": ("name", "org"),
}
_TAG_TYPES = {"skills", "languages", "interests"}
_MAX_TEXT = 4000


def sanitize_filename(name: str) -> str:
    name = re.sub(r"[^\w\s\-\.]", "", name or "", flags=re.U)
    name = re.sub(r"\s+", "_", name).strip("_")
    return name or "Document"


def _s(value: Any) -> str:
    return str(value or "").strip()[:_MAX_TEXT]


def _accent(cv: Dict[str, Any]) -> RGBColor:
    m = _HEX.match(_s((cv.get("style") or {}).get("accent")))
    return RGBColor.from_string(m.group(1).upper()) if m else RGBColor(0x2F, 0x5D, 0x8A)


def _setup(cv: Dict[str, Any]) -> Tuple[Document, RGBColor, float]:
    style = cv.get("style") or {}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_SIZES.get(style.get("page"), PAGE_SIZES["A4"])
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin = sec.bottom_margin = Mm(16)

    font_name = (FONTS.get(style.get("font")) or FONTS["inter"])["docx"]
    base = BASE_PT.get(style.get("size"), BASE_PT["m"])
    normal = doc.styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(base)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    return doc, _accent(cv), base


def _bottom_border(paragraph, color: RGBColor) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), str(color))
    borders.append(bottom)
    p_pr.append(borders)


def _right_tab(paragraph, doc: Document) -> None:
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)


def _header(doc: Document, cv: Dict[str, Any], accent: RGBColor, base: float) -> None:
    p = cv.get("personal") or {}
    name = " ".join(x for x in (_s(p.get("firstName")), _s(p.get("lastName"))) if x)
    para = doc.add_paragraph()
    run = para.add_run(name or "Your Name")
    run.bold = True
    run.font.size = Pt(base * 2.2)
    run.font.color.rgb = accent
    if _s(p.get("headline")):
        h = doc.add_paragraph()
        r = h.add_run(_s(p.get("headline")))
        r.font.size = Pt(base * 1.1)
    contact = [_s(p.get(k)) for k in ("email", "phone", "location", "linkedin", "website") if _s(p.get(k))]
    if contact:
        c = doc.add_paragraph(" | ".join(contact))
        c.paragraph_format.space_before = Pt(3)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def _heading(doc: Document, title: str, accent: RGBColor, base: float) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(4)
    run = para.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(base * 1.02)
    run.font.color.rgb = accent
    _bottom_border(para, accent)


def _description(doc: Document, text: str) -> None:
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return
    if len(lines) == 1 and not re.match(r"^[•\-*]", lines[0]):
        doc.add_paragraph(lines[0])
        return
    for ln in lines:
        doc.add_paragraph(re.sub(r"^[•\-*]\s*", "", ln), style="List Bullet")


def _section(doc: Document, sec: Dict[str, Any], accent: RGBColor, base: float) -> None:
    kind = _s(sec.get("type"))
    items: List[Dict[str, Any]] = [i for i in (sec.get("items") or []) if isinstance(i, dict)]
    items = [i for i in items if any(_s(v) for k, v in i.items() if k != "id")]
    if not items:
        return
    _heading(doc, _s(sec.get("title")) or kind.title(), accent, base)

    if kind == "summary":
        doc.add_paragraph(_s(items[0].get("text")))
        return
    if kind in _TAG_TYPES:
        parts = []
        for i in items:
            label = _s(i.get("name"))
            if label and _s(i.get("level")):
                label += f" ({_s(i.get('level'))})"
            if label:
                parts.append(label)
        doc.add_paragraph(", ".join(parts))
        return

    title_key, sub_key = _ENTRY_KEYS.get(kind, ("name", "org"))
    for n, item in enumerate(items):
        top = doc.add_paragraph()
        if n:
            top.paragraph_format.space_before = Pt(6)
        _right_tab(top, doc)
        top.add_run(_s(item.get(title_key))).bold = True
        dates = " – ".join(x for x in (_s(item.get("start")), _s(item.get("end"))) if x)
        if dates:
            top.add_run("\t" + dates)
        sub_bits = [x for x in (_s(item.get(sub_key)), _s(item.get("grade"))) if x]
        loc = _s(item.get("location"))
        if sub_bits or loc:
            sub = doc.add_paragraph()
            _right_tab(sub, doc)
            if sub_bits:
                sub.add_run(" · ".join(sub_bits)).italic = True
            if loc:
                sub.add_run("\t" + loc)
        _description(doc, _s(item.get("description")))


def _letter(doc: Document, cv: Dict[str, Any], accent: RGBColor, base: float) -> None:
    p = cv.get("personal") or {}
    letter = cv.get("letter") or {}
    name = " ".join(x for x in (_s(p.get("firstName")), _s(p.get("lastName"))) if x)
    doc.add_paragraph(date.today().strftime("%d %B %Y")).paragraph_format.space_after = Pt(10)
    for line in (_s(letter.get("recipient")), _s(letter.get("recipientTitle")), _s(letter.get("company")), _s(letter.get("address"))):
        if line:
            doc.add_paragraph(line)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)
    if _s(letter.get("role")):
        re_line = doc.add_paragraph()
        re_line.add_run(f"Re: {_s(letter.get('role'))}").bold = True
        re_line.paragraph_format.space_after = Pt(10)
    last = _s(letter.get("recipient")).split()[-1] if _s(letter.get("recipient")) else ""
    doc.add_paragraph(f"Dear {last or 'Hiring Manager'},").paragraph_format.space_after = Pt(8)
    for para in re.split(r"\n\s*\n", _s(letter.get("body"))):
        if para.strip():
            doc.add_paragraph(para.strip()).paragraph_format.space_after = Pt(8)
    doc.add_paragraph(_s(letter.get("signoff")) or "Kind regards,")
    doc.add_paragraph(name)


def build_docx(cv: Dict[str, Any], doc_kind: str = "cv") -> Tuple[bytes, str]:
    """Return (docx bytes, download filename) for the CV or its cover letter."""
    if not isinstance(cv, dict):
        raise ValueError("CV data must be an object")
    doc, accent, base = _setup(cv)
    _header(doc, cv, accent, base)
    if doc_kind == "letter":
        _letter(doc, cv, accent, base)
    else:
        for sec in cv.get("sections") or []:
            if isinstance(sec, dict) and not sec.get("hidden"):
                _section(doc, sec, accent, base)

    p = cv.get("personal") or {}
    who = "_".join(sanitize_filename(x) for x in (_s(p.get("firstName")), _s(p.get("lastName"))) if x) or "My"
    suffix = "Cover_Letter" if doc_kind == "letter" else "CV"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"{who}_{suffix}.docx"
