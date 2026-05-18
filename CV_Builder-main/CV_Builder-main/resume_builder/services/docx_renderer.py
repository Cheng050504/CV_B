"""DOCX rendering service.

Pure functions for placeholder replacement and experience-block cloning.
No Flask, no filesystem writes on the hot path — everything returns bytes
so the caller can stream them straight to the client.
"""

from __future__ import annotations

import io
import re
from copy import deepcopy
from datetime import datetime
from typing import Any, Dict, List

from docx import Document
from docx.text.paragraph import Paragraph


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def sanitize_filename(name: str) -> str:
    name = re.sub(r"[^\w\s\-\.]", "", name or "", flags=re.U)
    name = re.sub(r"\s+", "_", name).strip("_")
    return name or "Document"


def compose_state_country(state: str, country: str) -> str:
    """US → state only; otherwise 'state, country' (skip empty parts)."""
    s = (state or "").strip()
    c = (country or "").strip()
    if not s and not c:
        return ""
    if c and c.lower() in {"us", "usa", "united states", "united states of america"}:
        return s
    if s and c:
        return f"{s}, {c}"
    return s or c


def pack_repeating(form, prefix: str, fields: List[str]) -> List[Dict[str, str]]:
    """Collect repeating list-style form fields into a list of dicts.

    `form` is anything with a `getlist(name)` method — typically
    `flask.request.form`. Also accepts a plain dict (for tests).
    """
    def getlist(name: str) -> List[str]:
        if hasattr(form, "getlist"):
            return form.getlist(name)
        val = form.get(name, [])
        return val if isinstance(val, list) else [val]

    lists = {f: getlist(f"{prefix}_{f}[]") for f in fields}
    length = max((len(v) for v in lists.values()), default=0)
    items: List[Dict[str, str]] = []
    for i in range(length):
        item = {f: (lists[f][i].strip() if i < len(lists[f]) else "") for f in fields}
        if any(item.values()):
            items.append(item)
    return items


# ---------------------------------------------------------------------------
# Run-preserving placeholder replacement
# ---------------------------------------------------------------------------

def replace_in_runs_preserve(paragraph, mapping: Dict[str, str]) -> None:
    """Replace placeholders inside a paragraph without destroying run-level formatting.

    Word frequently splits a literal placeholder like ``[Company Name]`` across
    several ``<w:r>`` runs (e.g. after spell-check). Assigning to
    ``paragraph.text`` flattens all runs and loses bold/italic/font. This
    function walks the runs, builds a virtual "full text" string with an index
    map, finds longest-first non-overlapping matches, then rewrites the exact
    runs that contain each match, preserving formatting on everything else.
    """
    runs = paragraph.runs
    if not runs:
        return

    def _norm(s: Any) -> str:
        if s is None:
            return ""
        return str(s).replace("\r\n", "\n").replace("\r", "\n")

    norm_mapping = {k: _norm(v) for k, v in mapping.items()}

    run_texts = [r.text or "" for r in runs]
    index_map: List[tuple] = []
    full_pieces: List[str] = []
    for ri, t in enumerate(run_texts):
        full_pieces.append(t)
        index_map.extend((ri, oi) for oi in range(len(t)))
    full_text = "".join(full_pieces)
    if not full_text:
        return

    keys = sorted(norm_mapping.keys(), key=len, reverse=True)
    occupied = [False] * len(full_text)
    matches: List[tuple] = []

    for k in keys:
        if not k:
            continue
        klen = len(k)
        start = 0
        repl = norm_mapping[k]
        while True:
            pos = full_text.find(k, start)
            if pos == -1:
                break
            end = pos + klen
            if any(occupied[pos:end]):
                start = pos + 1
                continue
            matches.append((pos, end, repl))
            for i in range(pos, end):
                occupied[i] = True
            start = end

    if not matches:
        return

    for start, end, repl in sorted(matches, key=lambda x: x[0], reverse=True):
        ri, oi = index_map[start]
        rj, oj = index_map[end - 1]

        if ri == rj:
            t = runs[ri].text
            runs[ri].text = t[:oi] + repl + t[oj + 1:]
        else:
            t_first = runs[ri].text
            t_last = runs[rj].text
            runs[ri].text = t_first[:oi] + repl
            for m in range(ri + 1, rj):
                runs[m].text = ""
            runs[rj].text = t_last[oj + 1:]


def _replace_in_table(tbl, mapping: Dict[str, str]) -> None:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_in_runs_preserve(p, mapping)
            for nested_tbl in cell.tables:
                _replace_in_table(nested_tbl, mapping)


def replace_in_doc_preserve(doc: Document, mapping: Dict[str, str]) -> None:
    for p in doc.paragraphs:
        replace_in_runs_preserve(p, mapping)
    for tbl in doc.tables:
        _replace_in_table(tbl, mapping)


# ---------------------------------------------------------------------------
# Experience-block cloning (preserves DOCX formatting of the first block)
# ---------------------------------------------------------------------------

def _copy_paragraph_with_replacements(before_paragraph, src_paragraph, mapping):
    new_ctp = deepcopy(src_paragraph._element)
    before_paragraph._element.addprevious(new_ctp)
    new_p = Paragraph(new_ctp, before_paragraph._parent)
    replace_in_runs_preserve(new_p, mapping)
    return new_p


def _find_first_experience_block(doc: Document):
    paras = doc.paragraphs
    start = None
    for i, p in enumerate(paras):
        if "[Company Name]" in p.text:
            start = i
            break
    if start is None:
        return None, None
    end = len(paras)
    for j in range(start + 1, len(paras)):
        if "SKILLS, ACTIVITIES & INTERESTS" in paras[j].text:
            end = j
            break
    return start, end


def materialize_experiences(doc: Document, experiences: List[Dict[str, str]]) -> None:
    """Fill the first experience block, then clone it for each additional experience."""
    if not experiences:
        return
    start, end = _find_first_experience_block(doc)
    if start is None:
        return

    block_paras = doc.paragraphs[start:end]

    def exp_map(e: Dict[str, str]) -> Dict[str, str]:
        title = e.get("title", "")
        group = e.get("group", "")
        title_group = f"{title}, {group}" if group else title
        return {
            "[Company Name]": e.get("company", ""),
            "[City]": e.get("city", ""),
            "[State/Country]": compose_state_country(e.get("state", ""), e.get("country", "")),
            "[Position Title], [Group Name]": title_group,
            "[Start Date]": e.get("start", ""),
            "[End Date]": e.get("end", ""),
            "[Experience Description]": e.get("summary", ""),
        }

    block_paras_bak = deepcopy(block_paras)

    first_mapping = exp_map(experiences[0])
    for p in block_paras:
        replace_in_runs_preserve(p, first_mapping)

    # Bug 1 fix: end == len(paras) when boundary heading is absent — guard the index.
    if end >= len(doc.paragraphs):
        insert_before = doc.add_paragraph("")
    else:
        insert_before = doc.paragraphs[end]

    # Bug 2 fix: keep insert_before fixed so copies always land in document order.
    # Insert one blank separator only *between* additional blocks (not after the last).
    for i, e in enumerate(experiences[1:]):
        if i > 0:
            sep = insert_before.insert_paragraph_before("")
            sep.style = block_paras_bak[-1].style
        for src_p in block_paras_bak:
            _copy_paragraph_with_replacements(insert_before, src_p, exp_map(e))


# ---------------------------------------------------------------------------
# Top-level: extract form → rendered DOCX bytes for CV + Cover Letter
# ---------------------------------------------------------------------------

def _doc_to_bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_documents(form, cv_template_path: str, cl_template_path: str) -> Dict[str, Any]:
    """Build CV + Cover Letter DOCX files purely in memory from the submitted form.

    Returns a dict of bytes + metadata. No filesystem writes.
    """
    first_name = form.get("first_name", "").strip()
    last_name = form.get("last_name", "").strip()
    email = form.get("email", "").strip()
    phone = form.get("phone", "").strip()
    physical_address = form.get("physical_address", "").strip()
    full_name = f"{first_name} {last_name}".strip()

    education = pack_repeating(form, "ed", [
        "school", "city", "state", "country", "degree_type", "field", "start", "end",
        "gpa", "sat", "grade_system", "honors", "courses",
    ])
    ed = education[0] if education else {}

    experiences = pack_repeating(form, "e", [
        "company", "city", "state", "country", "title", "group", "start", "end", "summary",
    ])

    languages_fluent = form.get("languages", "").strip()
    languages_conversational = form.get("languages_secondary", "").strip()
    technical_skills = form.get("technical_skills", "").strip()
    certifications = form.get("certifications", "").strip()
    activities = form.get("activities", "").strip()
    interests = form.get("interests", "").strip()

    recruiter_name = form.get("recruiter_name", "").strip()
    recruiter_title = form.get("recruiter_title", "").strip()
    recruiter_last_name = recruiter_name.split()[-1] if recruiter_name else ""
    recruiter_salutation = form.get("recruiter_salutation", "").strip()
    company_name = form.get("company_name", "").strip()
    recruiter_address = form.get("recruiter_address", "").strip()

    cl_year = form.get("cl_year", "").strip()
    cl_school_name = ed.get("school", "")
    cl_major = ed.get("field", "")

    referral_source = form.get("referral_source", "").strip()
    firm_impression = form.get("firm_impression", "").strip()
    position_name = form.get("position_name", "").strip()
    past_experience = form.get("past_experience", "").strip()
    experience_theme = form.get("experience_theme", "").strip()
    gained_skills = form.get("gained_skills", "").strip()
    other_skills = form.get("other_skills", "").strip()
    project = form.get("project", "").strip()
    project_result = form.get("project_result", "").strip()
    background_summary = form.get("background_summary", "").strip()
    skill_summary = form.get("skill_summary", "").strip()
    firm_track_record = form.get("firm_track_record", "").strip()
    signature = form.get("signature", "").strip()

    # ---------------- CV ----------------
    cv_doc = Document(cv_template_path)
    materialize_experiences(cv_doc, experiences)
    cv_map = {
        "[Name]": full_name,
        "[Physical Address]": physical_address,
        "[Phone Number]": phone,
        "[Email Address]": email,
        "[University Name]": ed.get("school", ""),
        "[City]": ed.get("city", ""),
        "[State/Country]": compose_state_country(ed.get("state", ""), ed.get("country", "")),
        "[Arts/Science]": ed.get("degree_type", ""),
        "[Major]": ed.get("field", ""),
        "[Graduation Date]": ed.get("end", ""),
        "[GPA]": ed.get("gpa", ""),
        "[SAT]": ed.get("sat", ""),
        "[If you\u2019re outside the US, list grades under your system here instead]": ed.get("grade_system", ""),
        "[Honors]": ed.get("honors", ""),
        "[Economics / Accounting / Finance classes, anything business-related]": ed.get("courses", ""),
        "[Fluent]": languages_fluent,
        "[Conversational]": languages_conversational,
        "[List any programming languages \u2013 not MS Office/Excel]": technical_skills,
        "[Any extra courses or programs relevant to finance]": certifications,
        "[Student Clubs, Volunteer Work, Independent Activities]": activities,
        "[Keep this to 1-2 lines and be specific; do not go overboard]": interests,
    }
    replace_in_doc_preserve(cv_doc, cv_map)

    # ---------------- Cover Letter ----------------
    cl_doc = Document(cl_template_path)
    cl_map = {
        "[Your Name]": full_name,
        "[Your Address]": physical_address,
        "[Your Phone Number]": phone,
        "[Your Email Address]": email,
        "[Date]": datetime.today().strftime("%d %b %Y"),
        "[Name of Recruiter]": recruiter_name,
        "[Title]": recruiter_title,
        "[Name of Bank]": company_name,
        "[Recruiter\u2019s Address]": recruiter_address,
        "[Mr. / Ms.]": recruiter_salutation or "Mr.",
        "[Recruiter\u2019s Name]": recruiter_last_name,
        "[Year]": cl_year,
        "[School Name]": cl_school_name,
        "[Major]": cl_major,
        "[Friend / Contact at Firm / Presentation]": referral_source,
        "[Your Culture / Working Environment / Bank-Specific Info.]": firm_impression,
        "[Investment Banking Analyst / Associate]": position_name,
        "[Completed Internships In\u2026 / Worked Full-Time In\u2026]": past_experience,
        "[Working on Transactions / Leading Teams and Managing Projects / Performing Quantitative Analysis]": experience_theme,
        "[Go Into Anything Relevant to Banking, Such As Analytical / Leadership / Teamwork / Finance / Accounting]": gained_skills,
        "[Any Other Relevant Skills]": other_skills,
        "[High-Impact Project]": project,
        "[Describe Results]": project_result,
        "[Summarize Internships / Work Experience]": background_summary,
        "[Summarize Skills]": skill_summary,
        "[Position Name]": position_name,
        "[Transactions / Clients]": firm_track_record,
        "[Firm Name]": company_name,
        "[Phone Number]": phone,
        "[Email Address]": email,
        "[Signature]": signature,
    }
    replace_in_doc_preserve(cl_doc, cl_map)

    fname = sanitize_filename(first_name or "Firstname")
    lname = sanitize_filename(last_name or "Surname")

    return {
        "cv_docx_bytes": _doc_to_bytes(cv_doc),
        "cl_docx_bytes": _doc_to_bytes(cl_doc),
        "cv_filename_base": f"{fname}_{lname}_CV",
        "cl_filename_base": f"{fname}_{lname}_CoverLetter",
    }
